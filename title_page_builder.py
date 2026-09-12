"""
title_page_builder.py — Generator Halaman Judul (Title Page) JIP 2026
Template final v4:
  - Tabel penulis: 4 kolom (No | Nama | Aff | ORCID iD)
  - Seksi 4: CRediT + ORCID per penulis (dua blok terpisah)
  - Seksi 5: Konflik Kepentingan
  - Seksi 6: Sumber Dana
  - Seksi 7: Ucapan Terima Kasih
"""

import io
import copy
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

TEMPLATE_TITLE_PAGE = "Template_Title_Page_JIP_2026.docx"

CREDIT_ROLES = [
    "Conceptualization", "Data curation", "Formal analysis",
    "Funding acquisition", "Investigation", "Methodology",
    "Project administration", "Resources", "Software",
    "Supervision", "Validation", "Visualization",
    "Writing – original draft", "Writing – review & editing",
]


# ─────────────────────────────────────────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def _all_paras(doc):
    for p in doc.paragraphs:
        yield p
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def _find(doc, *kws):
    for p in _all_paras(doc):
        if any(k in p.text for k in kws):
            return p
    return None


def _fill2(para, label, value):
    runs = para.runs
    if len(runs) >= 2:
        runs[0].text = label
        runs[1].text = value
        for r in runs[2:]:
            r._r.getparent().remove(r._r)
    elif len(runs) == 1:
        runs[0].text = value
    else:
        r = para.add_run(value); r.font.italic = True


def _fill1(para, value):
    runs = para.runs
    if runs:
        runs[0].text = value
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        r = para.add_run(value); r.font.italic = True


def _make_bold_italic_p(ref_p_elem, label, value):
    """
    Buat elemen <w:p> baru (format disalin dari ref_p_elem):
    run[0] = label (bold), run[1] = value (italic).
    """
    new_p = copy.deepcopy(ref_p_elem)
    # Hapus semua run lama
    for r in new_p.findall(qn('w:r')):
        new_p.remove(r)

    def _run(text, bold=False, italic=False):
        r = OxmlElement('w:r')
        rpr = OxmlElement('w:rPr')
        if bold:
            rpr.append(OxmlElement('w:b'))
        if italic:
            rpr.append(OxmlElement('w:i'))
        r.append(rpr)
        t = OxmlElement('w:t')
        t.text = text
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        r.append(t)
        return r

    if label:
        new_p.append(_run(label, bold=True))
    new_p.append(_run(value, italic=True))
    return new_p


def _insert_after(ref_p_elem, new_p_elem):
    ref_p_elem.addnext(new_p_elem)


def _set_cell(cell, text, bold=False, italic=False, color="555555"):
    para = cell.paragraphs[0]
    for run in list(para.runs):
        run._r.getparent().remove(run._r)
    run = para.add_run(text)
    run.font.name = "Cambria"
    run.font.size = Pt(8)
    run.font.bold = bold
    run.font.italic = italic
    r, g, b = int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16)
    run.font.color.rgb = RGBColor(r, g, b)


def _clone_row(tr):
    new_tr = copy.deepcopy(tr)
    for tc in new_tr.findall(qn('w:tc')):
        for p_el in tc.findall(qn('w:p')):
            for r_el in p_el.findall(qn('w:r')):
                t_el = r_el.find(qn('w:t'))
                if t_el is not None:
                    t_el.text = ""
    return new_tr


# ─────────────────────────────────────────────────────────────────────────────
# FUNGSI UTAMA
# ─────────────────────────────────────────────────────────────────────────────

def bangun_title_page(data: dict, template_path: str = TEMPLATE_TITLE_PAGE) -> bytes:
    if not Path(template_path).exists():
        raise FileNotFoundError(f"Template tidak ditemukan: {template_path}")

    doc = Document(template_path)
    penulis_list = data.get("penulis", [])

    # ── 0. ID NASKAH OJS ────────────────────────────────────────────────────
    p = _find(doc, "ID Naskah OJS:")
    if p: _fill2(p, "ID Naskah OJS: ", data.get("ojs_id", ""))

    # ── 1. JUDUL ─────────────────────────────────────────────────────────────
    p = _find(doc, "Judul (Bahasa Indonesia):")
    if p: _fill2(p, "Judul (Bahasa Indonesia): ", data.get("judul_id", ""))
    p = _find(doc, "Title (English):")
    if p: _fill2(p, "Title (English): ", data.get("judul_en", ""))
    p = _find(doc, "Judul Singkat / Running Title:")
    if p: _fill2(p, "Judul Singkat / Running Title: ", data.get("running_title", ""))

    # ── 2. TABEL PENULIS (4 kolom) ───────────────────────────────────────────
    if penulis_list:
        tbl = next((t for t in doc.tables
                    if any("ORCID" in c.text for c in t.rows[0].cells)), None)
        if tbl:
            while len(tbl.rows) > 1:
                tbl.rows[-1]._tr.getparent().remove(tbl.rows[-1]._tr)
            tmpl_tr = tbl.rows[0]._tr
            for i, pd in enumerate(penulis_list, 1):
                new_tr = _clone_row(tmpl_tr)
                tbl._tbl.append(new_tr)
                cells = tbl.rows[-1].cells
                nama_lbl = pd.get("nama","") + ("  ★" if pd.get("is_corresp") else "")
                aff_lbl  = f"({pd.get('aff','')})" if pd.get("aff") else "—"
                _set_cell(cells[0], str(i)+".",    bold=True,   color="007BB8")
                _set_cell(cells[1], nama_lbl,      italic=True, color="555555")
                _set_cell(cells[2], aff_lbl,       bold=True,   color="007BB8")
                _set_cell(cells[3], pd.get("orcid","—") or "—", italic=True, color="555555")

    # ── KETERANGAN AFILIASI ───────────────────────────────────────────────────
    for huruf, teks in data.get("afiliasi", {}).items():
        label = f"({huruf})"
        for p in doc.paragraphs:
            if p.text.strip().startswith(label):
                runs = p.runs
                if len(runs) >= 2: runs[1].text = teks
                elif runs: runs[0].text = f"{label} {teks}"
                break

    # ── 3. PENULIS KORESPONDENSI ─────────────────────────────────────────────
    for label, key in [
        ("Nama Lengkap / Full Name: ",  "koresp_nama"),
        ("Afiliasi / Affiliation: ",    "koresp_afiliasi"),
        ("Alamat / Address: ",          "koresp_alamat"),
        ("Telepon / Phone: ",           "koresp_telepon"),
        ("E-mail: ",                    "koresp_email"),
        ("ORCID iD: ",                  "koresp_orcid"),
    ]:
        p = _find(doc, label.strip())
        if p: _fill2(p, label, data.get(key, ""))

    # ── 4. CRediT + ORCID per penulis ────────────────────────────────────────
    #
    # Cari 4 anchor di doc.paragraphs:
    #   A = paragraf instruksi "Nyatakan kontribusi..."          [30]
    #   B = paragraf "Peran CRediT / Available Roles:"           [34]
    #   C = paragraf pertama baris ORCID contoh (setelah B)      [35]
    #   D = paragraf spasi " " setelah baris ORCID terakhir      [38]
    #
    # Langkah:
    #  1. Hapus baris CRediT contoh (antara A+1 dan B-1)
    #  2. Sisipkan baris CRediT baru setelah A
    #  3. Hapus baris ORCID contoh (antara B+1 dan D-1)
    #  4. Sisipkan baris ORCID baru setelah B

    kontribusi = data.get("kontribusi", {})
    paras = doc.paragraphs

    idx_A = next((i for i,p in enumerate(paras) if "Nyatakan kontribusi" in p.text), None)
    idx_B = next((i for i,p in enumerate(paras) if "Peran CRediT / Available Roles" in p.text), None)
    idx_D = next((i for i,p in enumerate(paras)
                  if i > (idx_B or 0) and p.text.strip() == ""), None)

    if idx_A is not None and idx_B is not None:
        # 1. Hapus baris CRediT contoh
        credit_contoh = list(paras[idx_A+1 : idx_B])
        for p in credit_contoh:
            p._p.getparent().remove(p._p)

        # 2. Sisipkan baris CRediT baru (urutan normal: gunakan addnext secara terbalik)
        paras2  = doc.paragraphs
        p_A     = next(p for p in paras2 if "Nyatakan kontribusi" in p.text)
        ref_p   = p_A._p
        for nama, roles in reversed(list(kontribusi.items())):
            new_elem = _make_bold_italic_p(ref_p, nama + ": ", roles)
            _insert_after(ref_p, new_elem)

        # 3. Hapus baris ORCID contoh
        paras3  = doc.paragraphs
        idx_B3  = next((i for i,p in enumerate(paras3) if "Peran CRediT / Available Roles" in p.text), None)
        idx_D3  = next((i for i,p in enumerate(paras3)
                        if i > (idx_B3 or 0) and p.text.strip() == ""), None)
        if idx_B3 is not None:
            end3 = idx_D3 if idx_D3 else len(paras3)
            orcid_contoh = list(paras3[idx_B3+1 : end3])
            for p in orcid_contoh:
                p._p.getparent().remove(p._p)

        # 4. Sisipkan baris ORCID baru
        paras4 = doc.paragraphs
        p_B4   = next((p for p in paras4 if "Peran CRediT / Available Roles" in p.text), None)
        if p_B4:
            ref_p4 = p_B4._p
            for pd in reversed(penulis_list):
                new_elem4 = _make_bold_italic_p(ref_p4, pd.get("nama","") + ": ",
                                                 pd.get("orcid","—") or "—")
                _insert_after(ref_p4, new_elem4)

    # ── 5. KONFLIK KEPENTINGAN ────────────────────────────────────────────────
    # Template hanya berisi baris contoh (Tanpa konflik / Ada konflik).
    # Sisipkan satu paragraf aktif di bawah baris "Ada konflik..." lalu hapus
    # baris contoh agar dokumen rapi — ATAU cukup sisipkan saja.
    # Keputusan: sisipkan 1 paragraf aktif setelah "Ada konflik / Conflict exists:"
    konflik_val = data.get("konflik", "")
    p_ada = _find(doc, "Ada konflik / Conflict exists:")
    if p_ada and konflik_val:
        new_konflik = _make_bold_italic_p(p_ada._p, "", konflik_val)
        _insert_after(p_ada._p, new_konflik)

    # ── 6. SUMBER DANA ────────────────────────────────────────────────────────
    p_dana = _find(doc, "Penelitian ini didanai oleh Direktorat",
                   "This study was funded by the Directorate")
    if p_dana:
        _fill1(p_dana, data.get("pendanaan", ""))

    # ── 7. UCAPAN TERIMA KASIH ───────────────────────────────────────────────
    p_ucapan = _find(doc, "Penulis mengucapkan terima kasih kepada [pihak",
                     "The authors thank [parties")
    if p_ucapan:
        _fill1(p_ucapan, data.get("ucapan", ""))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def get_credit_roles():
    return CREDIT_ROLES
